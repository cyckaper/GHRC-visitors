#!/usr/bin/env python3
"""產生測試用的合成母簡報（不是中心的真簡報）。

用途：讓 deck CLI 與 slim-master.py 的測試有一份結構完整的 pptx：
英中成對段落、表格式流程頁、五張人像的組織頁、內嵌影片、超大圖片、共用媒體、講者備忘。

    python3 scripts/make-fixture.py test/fixtures/generated/master-fixture.pptx
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

out = sys.argv[1] if len(sys.argv) > 1 else "test/fixtures/generated/master-fixture.pptx"
os.makedirs(os.path.dirname(out) or ".", exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TITLE_ONLY = prs.slide_layouts[5]


def add_title(slide, text, zh=None, size=36):
    """title placeholder；英文主標，中文為第二段落。"""
    t = slide.shapes.title
    t.text_frame.text = text
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    if zh:
        p = t.text_frame.add_paragraph()
        r = p.add_run()
        r.text = zh
        r.font.size = Pt(size - 12)
    return t


def add_pairs(slide, pairs, left=1, top=2.2, width=11, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(4))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for en, zh in pairs:
        for txt, sz in ((en, size), (zh, size - 4)):
            if txt is None:
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(sz)
    return box


def png_bytes(w, h, color, noise=False):
    img = Image.new("RGB", (w, h), color)
    if noise:
        import random
        px = img.load()
        rnd = random.Random(1)
        for y in range(0, h, 2):
            for x in range(0, w, 2):
                px[x, y] = (rnd.randrange(256), rnd.randrange(256), rnd.randrange(256))
    b = io.BytesIO()
    img.save(b, format="PNG")
    return b.getvalue()


tmpdir = tempfile.mkdtemp()
portrait = os.path.join(tmpdir, "portrait.png")
with open(portrait, "wb") as f:
    f.write(png_bytes(300, 400, (120, 160, 130)))
big = os.path.join(tmpdir, "big.png")
with open(big, "wb") as f:
    f.write(png_bytes(4000, 3000, (200, 200, 220), noise=True))
poster = os.path.join(tmpdir, "poster.png")
with open(poster, "wb") as f:
    f.write(png_bytes(640, 360, (40, 40, 40)))

# 影片：有 ffmpeg 就做一支真的 1 秒 mp4，否則放假位元組（結構測試用）
video = os.path.join(tmpdir, "clip.mp4")
ffmpeg = shutil.which("ffmpeg") or next((os.path.join(d, n) for d in ["/opt/pw-browsers/ffmpeg-1011"] if os.path.isdir(d) for n in os.listdir(d) if n.startswith("ffmpeg")), None)
made = False
if ffmpeg:
    try:
        subprocess.run([ffmpeg, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=green:s=64x64:d=1", "-pix_fmt", "yuv420p", video], check=True, timeout=60)
        made = os.path.getsize(video) > 0
    except Exception as e:  # noqa: BLE001
        print("ffmpeg failed, using dummy bytes:", e)
if not made:
    with open(video, "wb") as f:
        f.write(b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 200000)

# 1 封面
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Welcome to the Green Health Research Center", "歡迎蒞臨綠色健康研究中心")
add_pairs(s, [("Visiting Organisation Name", "來訪單位名稱"), ("Guest Name, Title", "來賓姓名 職稱"), ("1 January 2026 · 2026年1月1日", None)], top=3.0)

# 2 今日流程（表格）
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Programme for Today", "今日流程")
rows = [("10:00 – 10:20", "Center overview and the laboratories", "中心簡介與研究室概覽", "01 – 12"), ("10:20 – 11:00", "Laboratory tour, Rooms 301 to 305", "研究室參訪 301–305", "—"), ("11:00 – 11:30", "Discussion", "座談", "—")]
tbl = s.shapes.add_table(len(rows) + 1, 4, Inches(1), Inches(2.2), Inches(11), Inches(3)).table
for j, h in enumerate(["Time", "Programme", "議程", "Pages"]):
    tbl.cell(0, j).text = h
for i, r in enumerate(rows, start=1):
    for j, v in enumerate(r):
        tbl.cell(i, j).text = v
s.notes_slide.notes_text_frame.text = "Presenter notes for the programme slide."

# 3 Contents
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Contents", "簡報架構")
add_pairs(s, [("01 Why Now", "為什麼是現在"), ("02 Core Proposition", "核心主張"), ("03 Four Laboratories", "四間研究室"), ("04 Research Outcomes", "研究成果")])

# 4 The Claim
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "The Claim", "核心宣稱")
add_pairs(s, [("Measure · Design · Prescribe", "量測 · 設計 · 處方")])

# 5 Organisation：五張人像 ＋ 名字
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Organisation", "組織架構")
names = [("Chun-Yen Chang", "張俊彥 301"), ("Bau-Show Lin", "林寶秀 302"), ("Hui-Mei Chen", "陳惠美 303"), ("Po-Ju Chang", "張伯茹 304"), ("Chia-Kuen Cheng", "鄭佳昆 305")]
for i, (en, zh) in enumerate(names):
    s.shapes.add_picture(portrait, Inches(0.8 + i * 2.5), Inches(2.2), width=Inches(1.5))
    add_pairs(s, [(en, zh)], left=0.6 + i * 2.5, top=4.4, width=2.3, size=14)

# 6 分隔頁
s = prs.slides.add_slide(BLANK)
add_pairs(s, [("01 Why Now", "為什麼是現在")], top=3, size=44)

# 7 影片頁
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Lab 301 — Measurement in Practice", "301 量測實況")
s.shapes.add_movie(video, Inches(3), Inches(2.2), Inches(7), Inches(3.9), poster_frame_image=poster, mime_type="video/mp4")
s.notes_slide.notes_text_frame.text = "Play the film."

# 8 超大圖片
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Lab 302 — Simulation Outputs", "302 模擬結果")
s.shapes.add_picture(big, Inches(2), Inches(2.0), width=Inches(9))

# 9 中文段落（共用 portrait 圖）
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Lab 303 — Landscape Simulation Lab", "303 景觀環境模擬室 · 負責人 陳惠美 Hui-Mei Chen")
add_pairs(s, [("VR head-mounted displays and eye tracking", "VR 頭戴顯示器與眼動追蹤"), ("360 VR content library", "360VR 教材庫")])
s.shapes.add_picture(portrait, Inches(10.5), Inches(2.2), width=Inches(1.5))

# 10 Thank you（共用 portrait 圖）
s = prs.slides.add_slide(TITLE_ONLY)
add_title(s, "Thank you", "謝謝")
s.shapes.add_picture(portrait, Inches(11), Inches(5.5), width=Inches(1.2))

prs.save(out)
shutil.rmtree(tmpdir, ignore_errors=True)
print(f"wrote {out} ({os.path.getsize(out) / 1e6:.1f} MB, {len(prs.slides)} slides, video={'real' if made else 'dummy'})")
